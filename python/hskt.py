import os
import streamlit as st
from docx import Document
import google.generativeai as genai

# Cấu hình giao diện Streamlit
st.set_page_config(page_title="Trợ lý Sư phạm HSKT", layout="centered")

st.markdown("<h2 style='text-align: center; color: #003366;'>📚 TRỢ LÝ AI SOẠN GIÁO ÁN HÒA NHẬP HSKT</h2>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center; color: gray;'>Hệ thống đọc hiểu giáo án đại trà, tự động điều chỉnh giảm tải và phân hóa chuẩn sư phạm.</p>", unsafe_allow_html=True)
st.markdown("---")

# --- HÀM TỰ ĐỘNG DÒ TÌM MODEL PHÙ HỢP VỚI API KEY ---
def get_working_model(api_key):
    genai.configure(api_key=api_key)
    try:
        # Lấy danh sách các model khả dụng từ tài khoản API của người dùng
        available_models = []
        for m in genai.list_models():
            if 'generateContent' in m.supported_generation_methods:
                available_models.append(m.name)
        
        # Ưu tiên tìm các dòng model flash hoặc pro
        for pref in ['gemini-1.5-flash', 'gemini-1.5-pro', 'gemini-pro', 'gemini-1.0-pro']:
            for am in available_models:
                if pref in am:
                    return genai.GenerativeModel(am)
                    
        # Nếu có model bất kỳ hỗ trợ generateContent
        if available_models:
            return genai.GenerativeModel(available_models[0])
            
    except Exception as e:
        print(f"Lỗi dò model: {e}")
    
    # Dự phòng mặc định nếu không list được
    return genai.GenerativeModel('gemini-1.5-flash')

# --- THANH BÊN (SIDEBAR) CẤU HÌNH ---
st.sidebar.header("⚙️ Cấu hình hệ thống")

gemini_key = st.sidebar.text_input("Nhập API Gemini:", type="password", help="Dán mã API Google Gemini của bạn vào đây")

if st.sidebar.button("🔍 Kiểm tra kết nối khóa"):
    if not gemini_key:
        st.sidebar.error("Vui lòng nhập Key trước khi kiểm tra!")
    else:
        try:
            test_model = get_working_model(gemini_key.strip())
            test_model.generate_content("Hi")
            st.sidebar.success("🎉 Khóa API hợp lệ và kết nối thành công!")
        except Exception as e:
            st.sidebar.error(f"❌ Key không hợp lệ hoặc lỗi kết nối: {e}")

st.sidebar.markdown("---")

disability_type = st.sidebar.selectbox(
    "Chọn loại bất lợi cho học sinh:",
    (
        "Khuyết tật trí tuệ / Nhận thức nhẹ (chậm tiếp thu, cần trực quan)",
        "Khuyết tật vận động / Chân tay (khó viết, hạn chế thao tác tay)",
        "Khuyết tật thị giác / Về mắt (kém mắt, khó nhìn chữ nhỏ)",
        "Khuyết tật sức khỏe / Thể chất (dễ mệt mỏi, sức bền kém)",
        "Khuyết tật thính giác / Nghe"
    )
)

teacher_name = st.text_input("Họ và tên Giáo viên thực hiện:", placeholder="Ví dụ: Nguyễn Văn A")
school_name = st.text_input("Tên Trường / Tổ chuyên môn:", placeholder="Ví dụ: Trường THPT Sào Bảy")

# --- KHU VỰC TẢI FILE LÊN ---
st.subheader("1. Tải lên file giáo án đại trà (.docx)")
uploaded_file = st.file_uploader("Chọn file Word giáo án của bạn:", type=["docx"])

file_content = ""
if uploaded_file is not None:
    try:
        doc_reader = Document(uploaded_file)
        file_content = "\n".join([p.text for p in doc_reader.paragraphs if p.text.strip()])
        st.success(f"Đã đọc thành công tệp: {uploaded_file.name} ({len(file_content)} ký tự)")
    except Exception as e:
        st.error(f"Lỗi đọc file Word: {e}")

# --- XỬ LÝ KHI BẤM NÚT ---
if st.button("🚀 Bắt đầu Phân tích & Viết Giáo án HSKT", use_container_width=True):
    if not gemini_key:
        st.error("Vui lòng nhập và kiểm tra Khóa API Gemini ở thanh bên trái!")
    elif not file_content:
        st.error("Vui lòng tải lên file Word giáo án đại trà trước khi tiếp tục!")
    else:
        try:
            # Tự động chọn model tương thích chuẩn xác từ API Key của thầy
            model = get_working_model(gemini_key.strip())

            prompt = f"""
            Bạn là một chuyên gia cao cấp về phương pháp giảng dạy giáo dục hòa nhập và xây dựng Kế hoạch bài dạy tại Việt Nam.
            Dựa vào nội dung giáo án đại trà được cung cấp bên dưới, hãy biên soạn lại thành một bản **Kế hoạch bài dạy hòa nhập hoàn chỉnh** dành cho học sinh khuyết tật dạng: [{disability_type}].

            YÊU CẦU NGHIỆP VỤ:
            1. Không chèn ép hay cắt cụt nội dung cơ bản của lớp đại trà.
            2. Viết lại một cách mượt mà, văn phong sư phạm chuẩn mực, tự nhiên.
            3. Phân hóa rõ ràng trong từng hoạt động: 
               - Mục tiêu: Có tiêu chí riêng đã giảm tải cho HSKT.
               - Tiến trình dạy học: Ghi chú rõ phương án hỗ trợ, thay thế bài tập phức tạp bằng các bài tập trực quan, ngắn gọn, phù hợp với đối tượng khuyết tật được chọn.

            THÔNG TIN GIÁO VIÊN:
            - Giáo viên thực hiện: {teacher_name if teacher_name else "........................................"}
            - Đơn vị: {school_name if school_name else "........................................"}

            NỘI DUNG GIÁO ÁN ĐẠI TRÀ GỐC:
            ---
            {file_content}
            ---
            """

            with st.spinner("🤖 AI đang tự động quét model và biên soạn giáo án hòa nhập... Vui lòng đợi trong giây lát!"):
                response = model.generate_content(prompt)
                edu_content = response.text

            # --- ĐÓNG GÓI THÀNH FILE WORD (.docx) MỚI ---
            doc = Document()
            
            p_header = doc.add_paragraph()
            p_header.add_run(f"Đơn vị: {school_name if school_name else 'Trường THPT Sào Bảy'}\n").bold = True
            p_header.add_run(f"Họ và tên giáo viên: {teacher_name if teacher_name else '____________________'}\n")
            p_header.add_run(f"Đối tượng hỗ trợ giáo dục hòa nhập: {disability_type}\n")
            doc.add_paragraph("------------------------------------------------------------------------------------------------------------------")

            for line in edu_content.split('\n'):
                if line.startswith("# "):
                    doc.add_heading(line.replace("# ", ""), level=1)
                elif line.startswith("## "):
                    doc.add_heading(line.replace("## ", ""), level=2)
                elif line.startswith("### "):
                    doc.add_heading(line.replace("### ", ""), level=3)
                else:
                    doc.add_paragraph(line)

            output_filename = "GiaoAn_HoaNhap_HSKT_Chuan.docx"
            doc.save(output_filename)

            st.success("🎉 Giáo án đã được biên soạn thành công tuyệt đối!")
            
            with open(output_filename, "rb") as f:
                st.download_button(
                    label="📥 Tải xuống File Word Giáo án HSKT (.docx)",
                    data=f,
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

        except Exception as e:
            st.error(f"Đã xảy ra lỗi kết nối hoặc xử lý API: {e}")